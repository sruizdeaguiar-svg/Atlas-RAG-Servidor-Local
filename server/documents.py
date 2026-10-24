import csv, io, json, subprocess, hashlib
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document
from .store import ROOT
TOOLS = ROOT / 'build' / 'DocumentTools'
TEXT_EXT={'.txt','.md','.csv','.tsv','.json','.log','.r','.py','.html','.xml','.yaml','.yml'}
SUPPORTED=TEXT_EXT|{'.pdf','.xlsx','.docx'}
MAX_FILE=40*1024*1024

def file_hash(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''): h.update(block)
    return h.hexdigest()

def text_read(path):
    data=path.read_bytes()
    for enc in ('utf-8-sig','utf-16','cp1252'):
        try: return data.decode(enc)
        except UnicodeError: pass
    raise ValueError('Unsupported text encoding')

def tables(path):
    if path.suffix.lower()=='.xlsx':
        wb=load_workbook(path,read_only=True,data_only=True)
        out={}
        try:
            for sheet in wb:
                values=[]
                for i,r in enumerate(sheet.iter_rows(values_only=True)):
                    if i>200000: raise ValueError('Sheet exceeds 200000 rows')
                    values.append(list(r))
                out[sheet.title]=values
        finally: wb.close()
        return out
    if path.suffix.lower() in {'.csv','.tsv'}:
        t=text_read(path)
        try: dialect=csv.Sniffer().sniff(t[:8192],delimiters=',;\t')
        except csv.Error: dialect=csv.excel_tab if path.suffix.lower()=='.tsv' else csv.excel
        values=list(csv.reader(io.StringIO(t),dialect))
        if len(values)>200001: raise ValueError('Table exceeds 200000 rows')
        return {'Data':values}
    raise ValueError('Select CSV, TSV or XLSX')

def extract(path):
    ext=path.suffix.lower()
    if path.stat().st_size>MAX_FILE: raise ValueError('File exceeds 40 MiB')
    if ext=='.pdf':
        reader=PdfReader(path)
        if len(reader.pages)>500: raise ValueError('PDF exceeds 500 pages')
        pages=[(f'page {i+1}',p.extract_text() or '') for i,p in enumerate(reader.pages)]
        if any(len(t.strip())<20 for _,t in pages):
            p=subprocess.run([str(TOOLS),'--ocr',str(path)],capture_output=True,text=True,timeout=180)
            if p.returncode: raise ValueError('OCR failed')
            pages=[(f"page {p['page']}"+(' (OCR)' if p['ocr'] else ''),p['text']) for p in json.loads(p.stdout)]
        return pages
    if ext in {'.csv','.tsv','.xlsx'}:
        out=[]
        for name,values in tables(path).items():
            if not values: continue
            headers=[str(x) if x is not None else '' for x in values[0]]
            for i in range(1,len(values),12):
                text='Columns: '+ ' | '.join(headers)+'\n'+'\n'.join(' | '.join('' if x is None else str(x) for x in r) for r in values[i:i+12])
                out.append((f'sheet {name}, rows {i+1}-{min(i+12,len(values))}',text))
            if len(values)==1: out.append((f'sheet {name}, row 1',' | '.join(headers)))
        return out
    if ext=='.docx':
        d=Document(path)
        return [('document','\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows]))]
    if ext in TEXT_EXT: return [('text',text_read(path))]
    raise ValueError('Unsupported format')

def chunks(sections):
    out=[]
    for locator,text in sections:
        for i in range(0,len(text),1000):
            part=text[max(0,i-120):i+1000].strip()
            if part: out.append((locator+f', fragment {i//1000+1}',part))
    if len(out)>10000: raise ValueError('Document exceeds 10000 chunks')
    return out

def apple_embed(texts,language):
    p=subprocess.run([str(TOOLS)],input=json.dumps({'language':language,'texts':texts}),capture_output=True,text=True,timeout=120)
    if p.returncode: raise ValueError('Apple sentence embedding unavailable for '+language)
    return json.loads(p.stdout)
