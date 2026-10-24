import zipfile
from pathlib import Path
import pytest
from server import documents

def test_office_expansion_limit_before_parser(tmp_path):
    path=tmp_path/'bomb.docx'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml',b'x'*(33*1024*1024))
    assert path.stat().st_size<100000
    with pytest.raises(ValueError,match='expanded size'):documents.extract(path)

def test_table_width_is_bounded(tmp_path):
    path=tmp_path/'wide.csv';path.write_text(','.join(['x']*257))
    with pytest.raises(ValueError,match='safe row, column'):documents.tables(path)

def test_isolated_csv_xlsx_docx_remain_readable(tmp_path):
    from openpyxl import Workbook
    from docx import Document
    csv=tmp_path/'data.csv';csv.write_text('x,y\n1,2\n2,4\n')
    assert documents.tables(csv)['Data']==[['x','y'],['1','2'],['2','4']]
    xlsx=tmp_path/'data.xlsx';wb=Workbook();wb.active.append(['x','y']);wb.active.append([1,2]);wb.save(xlsx)
    assert documents.tables(xlsx)['Sheet'][1]==[1,2]
    docx=tmp_path/'data.docx';doc=Document();doc.add_paragraph('FICTITIOUS ISOLATED QA');doc.save(docx)
    assert 'FICTITIOUS ISOLATED QA' in documents.extract(docx)[0][1]

def test_spreadsheet_entities_are_rejected(tmp_path):
    from openpyxl import Workbook
    path=tmp_path/'entities.xlsx';wb=Workbook();wb.active.append(['x']);wb.save(path)
    with zipfile.ZipFile(path) as archive:entries={n:archive.read(n) for n in archive.namelist()}
    entries['xl/workbook.xml']=b'<!DOCTYPE workbook [<!ENTITY attack "UNTRUSTED ENTITY">]>'+entries['xl/workbook.xml']
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for n,data in entries.items():archive.writestr(n,data)
    with pytest.raises(ValueError):documents.tables(path)
